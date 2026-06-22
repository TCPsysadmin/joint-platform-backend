from __future__ import annotations

import hashlib
import posixpath
import re
from dataclasses import dataclass
from io import BytesIO
from typing import Any
from uuid import UUID

from fastapi import UploadFile
from supabase import AsyncClient

from orchestrator.config import settings
from orchestrator.supabase_json import JsonDict, as_dict_list

_CHUNK_SIZE = 1024 * 1024
_SUMMARY_CHARS = 500
_TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".csv",
    ".json",
    ".jsonl",
    ".log",
    ".srt",
    ".vtt",
}
_TEXT_CONTENT_TYPES = {
    "application/json",
    "application/jsonl",
    "application/x-ndjson",
    "text/csv",
    "text/markdown",
}
_PDF_CONTENT_TYPES = {"application/pdf"}
_DOCX_CONTENT_TYPES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


class DocumentProcessingError(ValueError):
    """Raised when an upload cannot be converted into usable session context."""


@dataclass
class ProcessedDocument:
    filename: str
    content_type: str | None
    byte_size: int
    char_count: int
    sha256: str
    content_text: str
    summary: str
    metadata: dict[str, object]


def _safe_filename(filename: str | None) -> str:
    name = posixpath.basename((filename or "uploaded-document").strip().replace("\\", "/"))
    return name or "uploaded-document"


def _extension(filename: str) -> str:
    dot = filename.rfind(".")
    return filename[dot:].lower() if dot >= 0 else ""


def _is_supported_text_file(filename: str, content_type: str | None) -> bool:
    normalized_type = (content_type or "").split(";", 1)[0].strip().lower()
    return (
        normalized_type.startswith("text/")
        or normalized_type in _TEXT_CONTENT_TYPES
        or _extension(filename) in _TEXT_EXTENSIONS
    )


def _document_kind(filename: str, content_type: str | None) -> str | None:
    extension = _extension(filename)
    normalized_type = (content_type or "").split(";", 1)[0].strip().lower()
    if extension == ".pdf" or normalized_type in _PDF_CONTENT_TYPES:
        return "pdf"
    if extension == ".docx" or normalized_type in _DOCX_CONTENT_TYPES:
        return "docx"
    if _is_supported_text_file(filename, content_type):
        return "text"
    return None


async def _read_upload_bytes(file: UploadFile, *, max_bytes: int) -> bytes:
    chunks: list[bytes] = []
    total = 0
    while True:
        chunk = await file.read(_CHUNK_SIZE)
        if not chunk:
            break
        total += len(chunk)
        if total > max_bytes:
            raise DocumentProcessingError(
                f"Upload exceeds the {max_bytes} byte session document limit."
            )
        chunks.append(chunk)
    return b"".join(chunks)


def _decode_text(raw: bytes) -> str:
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DocumentProcessingError(
            "Only UTF-8 text-like documents are supported right now."
        ) from exc
    return text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")


def _extract_pdf_text(raw: bytes) -> str:
    try:
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(raw))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:
        raise DocumentProcessingError(f"Failed to extract text from PDF: {exc}") from exc
    return "\n\n".join(page.strip() for page in pages if page.strip())


def _extract_docx_text(raw: bytes) -> str:
    try:
        from docx import Document

        document = Document(BytesIO(raw))
    except Exception as exc:
        raise DocumentProcessingError(f"Failed to extract text from DOCX: {exc}") from exc

    parts: list[str] = []
    parts.extend(p.text for p in document.paragraphs if p.text.strip())
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_text(raw: bytes, *, kind: str) -> tuple[str, str]:
    if kind == "pdf":
        return _extract_pdf_text(raw), "pypdf"
    if kind == "docx":
        return _extract_docx_text(raw), "python-docx"
    return _decode_text(raw), "utf-8"


def _summarize_text(text: str) -> str:
    compact = re.sub(r"\s+", " ", text).strip()
    if len(compact) <= _SUMMARY_CHARS:
        return compact
    return compact[: _SUMMARY_CHARS - 1].rstrip() + "..."


async def process_upload_file(
    file: UploadFile,
    *,
    max_bytes: int | None = None,
) -> ProcessedDocument:
    filename = _safe_filename(file.filename)
    content_type = file.content_type
    kind = _document_kind(filename, content_type)
    if kind is None:
        raise DocumentProcessingError(
            "Unsupported document type. Upload a PDF, DOCX, or UTF-8 text-like "
            "file such as markdown, CSV, JSON, log, SRT, or VTT."
        )

    raw = await _read_upload_bytes(
        file,
        max_bytes=max_bytes or settings.session_document_max_upload_bytes,
    )
    text, extraction_method = _extract_text(raw, kind=kind)
    text = text.strip()
    if not text:
        raise DocumentProcessingError("Uploaded document is empty after text extraction.")

    return ProcessedDocument(
        filename=filename,
        content_type=content_type,
        byte_size=len(raw),
        char_count=len(text),
        sha256=hashlib.sha256(raw).hexdigest(),
        content_text=text,
        summary=_summarize_text(text),
        metadata={"extension": _extension(filename), "document_kind": kind, "extraction": extraction_method},
    )


async def create_session_document(
    svc: AsyncClient,
    *,
    session_id: str,
    user_id: UUID,
    client_id: UUID,
    document: ProcessedDocument,
) -> JsonDict:
    payload: dict[str, Any] = {
        "session_id": session_id,
        "user_id": str(user_id),
        "client_id": str(client_id),
        "filename": document.filename,
        "content_type": document.content_type,
        "byte_size": document.byte_size,
        "char_count": document.char_count,
        "sha256": document.sha256,
        "content_text": document.content_text,
        "summary": document.summary,
        "metadata": document.metadata,
        "status": "ready",
    }
    response = (
        await svc.table("chat_session_documents")
        .insert(payload)
        .execute()
    )
    rows = as_dict_list(response.data if response else None)
    if not rows:
        raise RuntimeError("Failed to persist uploaded session document")
    return rows[0]


async def list_session_documents(
    svc: AsyncClient,
    *,
    session_id: str,
    user_id: UUID,
    max_docs: int | None = None,
) -> list[JsonDict]:
    response = (
        await svc.table("chat_session_documents")
        .select(
            "doc_id, session_id, filename, content_type, byte_size, char_count, "
            "content_text, summary, metadata, created_at"
        )
        .eq("session_id", session_id)
        .eq("user_id", str(user_id))
        .eq("status", "ready")
        .order("created_at", desc=True)
        .limit(max_docs or settings.session_document_max_count)
        .execute()
    )
    return as_dict_list(response.data if response else None)


def build_document_context(
    documents: list[dict[str, Any]],
    *,
    max_chars: int | None = None,
) -> str | None:
    if not documents:
        return None

    budget = max_chars or settings.session_document_context_chars
    parts: list[str] = []
    used = 0

    for index, doc in enumerate(documents, start=1):
        filename = str(doc.get("filename") or f"Document {index}")
        content = str(doc.get("content_text") or doc.get("summary") or "").strip()
        if not content:
            continue

        header = f"### Uploaded Document {index}: {filename}\n"
        available = budget - used - len(header)
        if available <= 0:
            break
        excerpt = content[:available]
        parts.append(header + excerpt)
        used += len(header) + len(excerpt)
        if used >= budget:
            break

    if not parts:
        return None

    return (
        "The user uploaded these session-scoped documents. Use them as private "
        "conversation context when relevant; do not invent details beyond them.\n\n"
        + "\n\n".join(parts)
    )
